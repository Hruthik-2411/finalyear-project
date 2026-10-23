"""Generates a comprehensive Word document (.docx) detailing the FinSecure Security Architecture."""
import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_heading_styled(doc, text, level):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(6)
    for run in h.runs:
        run.font.name = 'Calibri'
        if level == 1:
            run.font.size = Pt(18)
            run.font.bold = True
            run.font.color.rgb = RGBColor(16, 44, 87) # Deep Navy
        elif level == 2:
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = RGBColor(28, 80, 140)
        elif level == 3:
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(50, 50, 50)
    return h

def add_callout(doc, title, text, bg_hex="F0F4F8", border_color="1C508C"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Left border
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    left = OxmlElement('w:left')
    left.set(qn('w:val'), 'single')
    left.set(qn('w:sz'), '24') # 3pt
    left.set(qn('w:space'), '0')
    left.set(qn('w:color'), border_color)
    tcBorders.append(left)
    for b in ['top', 'bottom', 'right']:
        node = OxmlElement(f'w:{b}')
        node.set(qn('w:val'), 'none')
        tcBorders.append(node)
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    r1 = p.add_run(f"NOTE: {title}\n")
    r1.bold = True
    r1.font.name = 'Calibri'
    r1.font.size = Pt(10.5)
    r1.font.color.rgb = RGBColor(16, 44, 87)
    
    r2 = p.add_run(text)
    r2.font.name = 'Calibri'
    r2.font.size = Pt(10)
    r2.font.color.rgb = RGBColor(40, 40, 40)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def format_table_headers(tbl, col_names, bg_hex="102C57"):
    hdr_cells = tbl.rows[0].cells
    for i, name in enumerate(col_names):
        hdr_cells[i].text = name
        set_cell_background(hdr_cells[i], bg_hex)
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=150, right=150)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.name = 'Calibri'
            run.font.bold = True
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(255, 255, 255)

def style_data_cells(tbl, col_widths=None):
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, row in enumerate(tbl.rows[1:]):
        bg = "F9FBFD" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, cell in enumerate(row.cells):
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            p = cell.paragraphs[0]
            for run in p.runs:
                run.font.name = 'Calibri'
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(40, 40, 40)
    if col_widths:
        for row in tbl.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = width

def generate_report():
    doc = Document()
    
    # Page setup
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    # Title Section
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    run_title = title_p.add_run("FinSecure: Trust-Aware Zero-Trust Framework")
    run_title.font.name = 'Calibri'
    run_title.font.size = Pt(24)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(16, 44, 87)
    
    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(16)
    run_sub = sub_p.add_run("Continuous Behavioral Security Architecture for Autonomous Multi-Agent Systems\nComprehensive Technical Specification & Implementation Guide")
    run_sub.font.name = 'Calibri'
    run_sub.font.size = Pt(12)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(100, 100, 100)

    # 1. Executive Summary
    add_heading_styled(doc, "1. Executive Summary & Design Philosophy", 1)
    doc.add_paragraph(
        "Autonomous Multi-Agent Systems (MAS) present unique cybersecurity vulnerabilities that traditional perimeter defense "
        "and point-of-login authentication cannot address. When Large Language Model (LLM) agents collaborate, a single "
        "compromised agent or poisoned context (e.g., indirect prompt injection hidden in banking memos or policies) can trigger "
        "cascading failures, unauthorized privilege escalation, confused deputy attacks, and irreversible financial damage.\n\n"
        "FinSecure establishes a Trust-Aware Zero-Trust architecture enforcing Continuous Behavioral Security across five specialized "
        "banking agents (Orchestrator, Planner, Researcher, Executor, Auditor). The fundamental axiom of FinSecure is Complete Mediation: "
        "Agents never communicate directly with each other, never hold direct database connections, and cannot invoke external APIs "
        "independently. Instead, every single interaction is sealed into a cryptographically signed envelope and routed through the "
        "Zero-Trust Security Gateway."
    )

    add_callout(
        doc,
        "Zero-Trust Complete Mediation Axiom",
        "No entity inside the multi-agent mesh is trusted by default, regardless of role or network locality. "
        "Every inter-agent message and every tool invocation is verified at runtime against identity credentials, user-intent scope, "
        "declarative policies, and behavioral trust scores before execution."
    )

    # 2. System Architecture Decomposition
    add_heading_styled(doc, "2. System Architecture Decomposition", 1)
    doc.add_paragraph(
        "The architecture corresponds directly to the four macro-layers established in the formal methodology:"
    )
    
    doc.add_paragraph(
        "• Layer 1: Input Layer (Enterprise User):\n"
        "  Accepts authenticated staff prompts and banking commands. Receives the finalized, auditor-verified output.\n\n"
        "• Layer 2: LangGraph Multi-Agent Orchestration Layer:\n"
        "  Constructed using LangGraph state graphs. Houses the five specialized banking agents:\n"
        "  - Orchestrator: Ingests requests, establishes intent tokens, and delivers final answers.\n"
        "  - Planner: Deconstructs requests into ordered, granular operational steps.\n"
        "  - Researcher: Connects to read-only tools, customer document vaults, and policy stores.\n"
        "  - Executor: Connects to operational write tools, account balances, and database mutations.\n"
        "  - Auditor: Evaluates the execution trace against banking policies and flags anomalies.\n\n"
        "• Layer 3: Security Proxy Block (Zero-Trust Gateway):\n"
        "  The single choke-point intercepting all inter-agent messages and tool calls. Houses the Synchronous Critical Path "
        "and Asynchronous Analytics Track, Core Security Engines, and the Adaptive Decision Engine.\n\n"
        "• Layer 4: Audit & Compliance Layer:\n"
        "  Cryptographic Merkle ledger committing every event hash into tamper-evident audit trees with externally signed roots."
    )

    # 3. Synchronous Critical Path vs Asynchronous Analytics Track
    add_heading_styled(doc, "3. Synchronous Critical Path vs. Asynchronous Analytics Track", 1)
    doc.add_paragraph(
        "A central architectural challenge in securing autonomous multi-agent systems is balancing rigorous, zero-bypass security "
        "enforcement with acceptable system throughput and latency. If every analytical metric, anomaly scan, and ledger root calculation "
        "ran inline before every single message, the multi-agent graph would grind to a halt. Conversely, if security checks were performed "
        "asynchronously in the background, an adversary could execute an irreversible transaction before the block verdict arrived.\n\n"
        "FinSecure resolves this via a Dual-Track Gateway Pipeline: dividing security processing strictly into a Synchronous Critical Path "
        "(inline blocking enforcement) and an Asynchronous Analytics Track (parallel non-blocking intelligence)."
    )

    # Table of Dual Tracks
    tbl_tracks = doc.add_table(rows=1, cols=4)
    format_table_headers(tbl_tracks, ["Track", "Execution Model", "Component Modules", "System Responsibility"])
    
    track_data = [
        ("Synchronous Critical Path", "Blocking (Inline)\n(Request held until verdict)", 
         "• Identity Cache Validator\n• Signed Intent Scope Verifier\n• Provenance DAG Walker\n• Prompt Shield (Regex + LPG2)\n• OPA Policy Engine (RBAC/ABAC)\n• Burst Limiter\n• Adaptive Decision Engine", 
         "Complete Mediation: Guarantees no tool runs and no message is delivered unless proven authentic, in-scope, policy-compliant, and injection-free."),
        ("Asynchronous Analytics Track", "Non-Blocking (Parallel)\n(Dedicated worker thread)", 
         "• Behavioral Anomaly Detector\n• Trust Score Reward Engine\n• Anomaly Observation Logger\n• Merkle Batch Root Sealer\n• WebSocket Event Broadcaster", 
         "Continuous Behavioral Intelligence: Calculates statistical anomalies, updates long-term trust scores, writes tamper-evident Merkle roots, and streams real-time telemetry to the SOC dashboard.")
    ]
    for row in track_data:
        r = tbl_tracks.add_row()
        for idx, val in enumerate(row):
            r.cells[idx].text = val
    style_data_cells(tbl_tracks, [Inches(1.5), Inches(1.5), Inches(2.2), Inches(2.3)])

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    add_heading_styled(doc, "3.1 How the Synchronous Critical Path is Implemented", 2)
    doc.add_paragraph(
        "Implemented in backend/gateway/gateway.py inside the _process(env, session) method. When an agent emits an envelope "
        "(via _send() or _tool()), the workflow thread enters _process() and blocks:\n\n"
        "1. Identity & Signature Verification: The Gateway extracts env.cert_pem, checks it against the local CA, verifies the "
        "RSA-PSS/SHA-256 signature, and checks the revoked certificate table in SQLite. If invalid, it immediately blocks (Rule 1).\n"
        "2. Provenance Chain & Intent Scope Verification: The Gateway verifies that the message has a valid parent lineage leading back "
        "to the signed Intent Token. If it is a tool call, intent.check_tool_scope() checks whether the target tool and arguments match "
        "the user's original request scope (Rule 2).\n"
        "3. OPA Policy Engine Evaluation: Translates the action, agent role, and target resources into an OPA query input. Evaluates "
        "RBAC rules (allowed edges, allowed tools) and ABAC rules (limits, account active status, PII leakage) (Rule 3).\n"
        "4. Prompt Shield Inspection: Evaluates incoming messages and tool results through regex heuristic patterns and the Groq-hosted "
        "Llama Prompt Guard 2 model (86M). If high risk is detected, malicious payloads are redacted or blocked (Rule 4).\n"
        "5. Burst Limiter Sliding Window: Verifies that the agent has not exceeded read limits (max 20 per 60s) or write limits (max 5 per 60s) "
        "or amount velocity limits (Rs 3,00,000 per 300s) (Rule 6).\n"
        "6. Adaptive Decision Engine: Synthesizes all check results into a definitive verdict (APPROVE, RESTRICT, ESCALATE, BLOCK, ISOLATE). "
        "Only if the verdict is APPROVE is the envelope delivered or the Python banking tool invoked."
    )

    add_heading_styled(doc, "3.2 How the Asynchronous Analytics Track is Implemented", 2)
    doc.add_paragraph(
        "Implemented in backend/gateway/gateway.py lines 58-116. During Gateway initialization, a thread-safe Queue and daemon thread "
        "are instantiated:\n\n"
        "   self._q = queue.Queue()\n"
        "   self._worker = threading.Thread(target=self._async_loop, daemon=True)\n"
        "   self._worker.start()\n\n"
        "When an action successfully passes the synchronous critical path, the Gateway returns the result immediately to the requesting agent, "
        "and enqueues the analytical tasks onto self._q:\n\n"
        "• Asynchronous Trust Rewards: Calling _async_good(agent, weight_key, amount) executes on the background worker thread. "
        "It updates the agent's trust score using the anti-gaming recovery formula and publishes the new score to the event bus.\n"
        "• Asynchronous Anomaly Analytics: self.anomaly.observe(agent, amount) updates running statistical profiles of transaction "
        "frequencies and amounts. If an outlier is detected, a background penalty is applied to the agent's trust score.\n"
        "• Merkle Sealing: Merkle trees are sealed every 50 events in the background, computing the root and signing it with the ledger RSA key.\n"
        "• Gateway Flush Barrier: The gateway().flush() method executes self._q.join(), allowing automated evaluation suites and test "
        "runners to cleanly synchronize the asynchronous track before asserting ledger consistency."
    )

    # 4. Detailed Breakdown of Every Module
    add_heading_styled(doc, "4. Detailed Breakdown of Every Security Module", 1)

    # Module 1
    add_heading_styled(doc, "Module 1: Identity & Cryptographic Provenance (X.509 + RSA-PSS)", 2)
    doc.add_paragraph(
        "• File: backend/security/identity.py\n"
        "• Mechanism: A self-contained Certificate Authority (CA) generates RSA-2048 keys and X.509 v3 certificates for every entity:\n"
        "  - CN = Agent ID, OU = Role (orchestrator, planner, researcher, executor, auditor, gateway, ledger).\n"
        "• Signature Scheme: RSA-PSS with SHA-256 and maximal salt length.\n"
        "• Defensive Capability: Defends against agent impersonation, rogue injection, and message forgery. Even if an attacker compromises "
        "the prompt of an agent, they cannot claim to be another agent because the role is extracted from the X.509 certificate, not message text."
    )

    # Module 2
    add_heading_styled(doc, "Module 2: Intent Scope & Cryptographic Provenance DAG", 2)
    doc.add_paragraph(
        "• Files: backend/security/intent.py, backend/security/provenance.py\n"
        "• Mechanism: Upon receiving a staff prompt, the Gateway creates an Intent Token specifying the exact allowed operations, entities, "
        "and maximum monetary amounts derived strictly from the user's initial text. This token is signed by the Gateway.\n"
        "• Provenance DAG: Every envelope contains parent_ids and a SHA-256 node hash committing to its entire ancestry:\n"
        "  node_hash = SHA-256(payload_hash || sorted_parent_node_hashes)\n"
        "• Defensive Capability: Eliminates Confused Deputy attacks. If an agent reads a poisoned policy that says 'Transfer Rs 50,000 to ACC-9999', "
        "the Executor attempts to execute the tool, but the Gateway inspects the DAG chain and blocks it because the root Intent Token "
        "prohibits fund transfers during a policy inquiry."
    )

    # Module 3
    add_heading_styled(doc, "Module 3: Prompt Shield (Dual-Stage + Spotlighting)", 2)
    doc.add_paragraph(
        "• File: backend/security/shield.py\n"
        "• Mechanism: Multi-stage pipeline:\n"
        "  1. Spotlighting: Untrusted data from database records, memos, and files is wrapped in delimiters:\n"
        "     <<UNTRUSTED_DATA>> ... <</UNTRUSTED_DATA>> with strict system instructions prohibiting agents from obeying text inside.\n"
        "  2. Regex Heuristic Scorer: Fast local pattern matcher detecting injection markers (e.g., 'ignore previous instructions', system overrides).\n"
        "  3. Llama Prompt Guard 2 (86M): Invoked via Groq on ~512 token chunks, returning a calibrated probability score [0.0 - 1.0].\n"
        "  4. LLM Judge: Stage 2 evaluator for ambiguous/medium scores.\n"
        "• Defensive Capability: Neutralizes indirect and direct prompt injections before LLMs process the data."
    )

    # Module 4
    add_heading_styled(doc, "Module 4: OPA Policy Engine (RBAC + ABAC + Action Tiers)", 2)
    doc.add_paragraph(
        "• Files: backend/security/policy.py, policies/finsecure.rego, backend/tools/registry.py\n"
        "• Mechanism: Enforces declarative Open Policy Agent (Rego) rules backed by a local Python evaluator for zero-latency fallbacks.\n"
        "• Role-Based Access Control (RBAC): Strict communication graph (e.g., Orchestrator -> Planner -> Researcher -> Executor -> Auditor). "
        "Researcher may only call read tools; Executor may only call write tools; Planner and Orchestrator have no tool access.\n"
        "• Attribute-Based Access Control (ABAC): Enforces single transfer limits (Rs 1,00,000), daily account limits (Rs 2,00,000), active account "
        "status verification, minimum credit scores for loans (>= 650), and strict PII prevention (prohibiting emails/phones in fraud notes).\n"
        "• Action Tiers:\n"
        "  - Tier 0: Read operations (unrestricted subject to RBAC).\n"
        "  - Tier 1: Reversible mutations (e.g., freeze_account, add_case_note) -> executed immediately if checks pass.\n"
        "  - Tier 2: Irreversible mutations (e.g., transfer_funds, unfreeze_account, update_loan_status) -> blocked or escalated to human review."
    )

    # Module 5 & 6
    add_heading_styled(doc, "Module 5 & 6: Burst Limiter & Behavioral Trust Engine", 2)
    doc.add_paragraph(
        "• File: backend/security/behaviour.py\n"
        "• Burst Limiter: Enforces sliding-window rate limits (max 20 reads/60s, max 5 writes/60s) and financial velocity limits (max Rs 3,00,000/300s). "
        "Tripping the burst limiter triggers a 120s cooldown and RESTRICT verdict.\n\n"
        "• Behavioral Trust Engine (Reputation Laundering Defense):\n"
        "  Standard systems use linear moving averages where an attacker can execute 40 harmless micro-actions to quickly recover trust. "
        "FinSecure implements a Non-Linear Asymmetric Trust Formula:\n\n"
        "  1. Violation Penalty (Multiplicative Decay):\n"
        "     T <- T * (1 - s),   where s is violation severity (0.2 to 1.0).\n\n"
        "  2. Good Action Reward (Anti-Gaming Dampened Recovery):\n"
        "     T <- min( T + [ eta0 / (1 + n_W(t)) ] * w_a * (1 - T),   C(t) )\n"
        "     - eta0 / (1 + n_W(t)) is the anti-gaming coefficient: packing multiple actions in window W reduces their reward.\n"
        "     - w_a is the action weight (read=0.01, tier-1=0.03, tier-2=0.05).\n\n"
        "  3. Time-Bounded Recovery Ceiling C(t):\n"
        "     C(t) = T_v + (1 - T_v) * [ 1 - exp(-(t - t_v) / tau) ]\n"
        "     - Limits trust recovery strictly by real elapsed time (tau = 1800s / 30 mins), making rapid reputation laundering mathematically impossible."
    )

    # Module 7
    add_heading_styled(doc, "Module 7: Adaptive Decision Engine", 2)
    doc.add_paragraph(
        "• File: backend/security/decision.py\n"
        "• Mechanism: Evaluates all synchronous checks in a deterministic, prioritized sequence (Rules 1 to 9).\n"
        "• Verdict Taxonomy:\n"
        "  - APPROVE: All checks passed; message delivered or tool executed.\n"
        "  - RESTRICT: Burst rate limits exceeded; action throttled.\n"
        "  - ESCALATE: Tier-2 action with borderline trust (0.40 - 0.70) or tainted context held for human approval.\n"
        "  - BLOCK: Security violation (bad signature, broken chain, out-of-scope, OPA deny, high-risk injection).\n"
        "  - ISOLATE: Trust score dropped below 0.20 or 3 repeated blocks; agent certificate revoked and quarantined."
    )

    # Module 8
    add_heading_styled(doc, "Module 8: Tamper-Evident Merkle Audit Ledger", 2)
    doc.add_paragraph(
        "• Files: backend/security/audit.py, backend/security/report.py\n"
        "• Mechanism: Every decision, message, and state mutation is logged as an AuditEvent. The leaf hash is computed as:\n"
        "  leaf_hash = SHA-256(canonical_json(seq, timestamp, event_type, actor, payload))\n"
        "• External Root Sealing: Every 50 events, a Merkle tree is computed over the batch. The root hash is signed by the Ledger's RSA key "
        "and appended to data/ledger_roots.jsonl outside SQLite.\n"
        "• Integrity Verification: The verify() algorithm recomputes all leaves from database rows, rebuilds the Merkle trees, and validates "
        "them against the signed roots file. If an insider or attacker modifies, inserts, or deletes a row in SQLite, verify() pinpoints "
        "the exact altered sequence number."
    )

    # 5. Attack Surface Coverage
    add_heading_styled(doc, "5. Attack Surface Coverage & Threat Mitigation", 1)
    
    tbl_threats = doc.add_table(rows=1, cols=3)
    format_table_headers(tbl_threats, ["Threat Class", "Attack Scenario", "FinSecure Defense Mechanism"])
    
    threat_data = [
        ("Confused Deputy", "An agent reads a poisoned policy or memo instructing it to execute an unauthorized transfer.", 
         "Signed Intent Scope & Provenance DAG: The root Intent Token defines allowed operations. Out-of-scope actions trigger Rule 2 BLOCK."),
        ("Reputation Laundering", "An adversary compromises an agent, suffers a penalty, then spams 40 benign reads to artificially restore trust.", 
         "Anti-Gaming eta & Recovery Ceiling C(t): Recovery is throttled by elapsed time tau (30 mins). Packing actions yields diminishing returns."),
        ("Post-Compromise Behavior", "A compromised agent attempts rapid data exfiltration, database scraping, or account draining.", 
         "Burst Limiter & Anomaly Detector: Sliding-window quotas throttle harvesting; three repeated blocks trigger automatic ISOLATION."),
        ("Audit Tampering", "An insider or malicious admin alters audit records in SQLite to cover illicit transactions.", 
         "Signed Merkle Ledger: Roots are cryptographically sealed outside SQLite. Any DB row edit invalidates the Merkle inclusion proof.")
    ]
    for row in threat_data:
        r = tbl_threats.add_row()
        for idx, val in enumerate(row):
            r.cells[idx].text = val
    style_data_cells(tbl_threats, [Inches(1.8), Inches(2.6), Inches(3.1)])

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 6. Conclusion
    add_heading_styled(doc, "6. Implementation Summary & Deliverables", 1)
    doc.add_paragraph(
        "The FinSecure architecture provides end-to-end continuous mediation across autonomous LLM multi-agent operations. "
        "By enforcing strict cryptographic identity, provenance verification, declarative policy enforcement, and non-linear "
        "behavioral trust dynamics, FinSecure guarantees that autonomous agents operate strictly within the bounds of authorized human intent."
    )

    # Output file
    output_path = os.path.abspath("FinSecure_Security_Architecture_Report.docx")
    doc.save(output_path)
    print(f"Document successfully created at: {output_path}")

if __name__ == "__main__":
    generate_report()
